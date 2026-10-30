"""智算组网配置清单（xlsx）行读取适配器。

openpyxl 属于第三方 I/O 框架，**不得**进入领域层（import-linter L5）。
本模块只负责把 xlsx 字节流转成「首 sheet 的字符串行矩阵」，
后续业务解析（类型前向填充、表头/合计跳过、BomLine 构造）在领域层
`ai_roi.domain.calc.bom_estimate.parse_network_bom_rows` 完成。
"""

from __future__ import annotations

from io import BytesIO
from typing import List

from openpyxl import load_workbook


def read_bom_rows(data: bytes) -> List[List[str]]:
    """读取 xlsx 第一个 sheet，返回字符串行矩阵（None → 空串，去首尾空白）。"""
    wb = load_workbook(BytesIO(data), data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows: List[List[str]] = []
    for row in ws.iter_rows(values_only=True):
        rows.append(["" if cell is None else str(cell).strip() for cell in row])
    return rows


def parse_network_bom(data: bytes):
    """便捷入口：xlsx 字节 → 领域 `BomSpec`（行读取 + 领域解析）。"""
    from ai_roi.domain.calc.bom_estimate import parse_network_bom_rows

    return parse_network_bom_rows(read_bom_rows(data))
